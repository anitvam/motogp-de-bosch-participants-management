"""Printable starting grids in the layout of the 2025 workbook."""
from __future__ import annotations

import io
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.pagebreak import Break

from .classes import CLASS_TITLES
from .models import GridRow

LOGO = Path(__file__).resolve().parents[2] / "assets" / "logo.png"
FONT = "Century Gothic"
COLUMN_WIDTHS = {"A": 10.71, "B": 23.29, "C": 22.0, "D": 54.29, "E": 22.29}
HEADERS = ["n°", "NOME", "COGNOME", "MOTO", "ANNO DI \nPRODUZIONE"]

DAYS = ["Lunedì", "Martedì", "Mercoledì", "Giovedì", "Venerdì", "Sabato", "Domenica"]
MONTHS = ["Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno", "Luglio",
          "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre"]


def italian_date(d: date | None) -> str:
    """date(2025, 9, 14) -> 'Domenica 14 Settembre 2025'."""
    if d is None:
        return ""
    return f"{DAYS[d.weekday()]} {d.day} {MONTHS[d.month - 1]} {d.year}"


def build_workbook(
    grids: dict[str, list[list[GridRow]]],
    date_text: str,
    max_per_grid: int = 8,
    logo: Path | None = LOGO,
    draft_issues: dict[int, str] | None = None,
) -> Workbook:
    """Starting grids in the 2025 layout. With `draft_issues` (bike_id -> problems) the sheets are
    marked as a draft and get an extra DA VERIFICARE column."""
    draft = draft_issues is not None
    red = Font(name=FONT, size=14, bold=True, color="C00000")
    issue_font = Font(name=FONT, size=11, bold=True, color="C00000")
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")

    for class_name, all_turns in grids.items():
        turns = [t for t in all_turns if t]
        if not turns:
            continue
        ws = wb.create_sheet(class_name)
        for col, width in COLUMN_WIDTHS.items():
            ws.column_dimensions[col].width = width
        if draft:
            ws.column_dimensions["F"].width = 52
        # Each turn is one printed page: 4 title rows, header, the grid rows and one spacer.
        block_rows = 5 + max_per_grid + 1
        for t, rows in enumerate(turns):
            top = 1 + t * block_rows
            for r in range(top, top + block_rows):
                ws.row_dimensions[r].height = 21
            ws.merge_cells(start_row=top, start_column=4, end_row=top, end_column=5)
            cell = ws.cell(top, 4, date_text)
            cell.font, cell.alignment = Font(name=FONT, size=18, bold=True), right
            cell = ws.cell(top + 1, 5, f"BATTERIA DELLE MOTO {CLASS_TITLES.get(class_name, class_name)}")
            cell.font, cell.alignment = Font(name=FONT, size=18), right
            if len(turns) > 1:
                cell = ws.cell(top + 2, 5, f"{t + 1}° TURNO")
                cell.font, cell.alignment = Font(name=FONT, size=18, bold=True), right
            if draft:
                ws.merge_cells(start_row=top + 3, start_column=4, end_row=top + 3, end_column=6)
                cell = ws.cell(top + 3, 4, "BOZZA DA VERIFICARE – non è la batteria ufficiale")
                cell.font, cell.alignment = red, right
            if logo and Path(logo).exists():
                img = Image(str(logo))
                img.width, img.height = 237, 97
                ws.add_image(img, f"A{top}")

            header_row = top + 4
            ws.row_dimensions[header_row].height = 46
            for c, text in enumerate(HEADERS + (["DA VERIFICARE"] if draft else []), start=1):
                cell = ws.cell(header_row, c, text)
                cell.font, cell.alignment, cell.border = Font(name=FONT, size=18, bold=True), center, border
            for i in range(max_per_grid):
                r = header_row + 1 + i
                ws.row_dimensions[r].height = 27
                row = rows[i] if i < len(rows) else None
                values = (
                    [row.race_number, row.first_name, row.last_name, row.bike, row.year]
                    if row else [None] * 5
                )
                for c, value in enumerate(values, start=1):
                    cell = ws.cell(r, c, value)
                    cell.font, cell.alignment, cell.border = Font(name=FONT, size=18), center, border
                if draft:
                    cell = ws.cell(r, 6, draft_issues.get(row.bike_id, "ok") if row else None)
                    cell.alignment, cell.border = center, border
                    cell.font = issue_font if row and row.bike_id in draft_issues else Font(name=FONT, size=14, color="2E7D32")
            if t < len(turns) - 1:
                ws.row_breaks.append(Break(id=top + block_rows - 1))

        last = len(turns) * block_rows
        ws.print_area = f"A1:{'F' if draft else 'E'}{last}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.24
        ws.print_options.horizontalCentered = True

    if not wb.worksheets:
        wb.create_sheet("Nessuna batteria")
    return wb


def to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
