"""Full data export for people to read, print and fill in by hand.

Sheet 'Iscritti': each participant written once (cells merged over their bikes), with every
bike's race number right next to it, and an empty PAGATO box to tick by hand.
Sheet 'Moto per numero': quick lookup of a bike by its race number.
"""
from __future__ import annotations

from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .classes import CLASS_TITLES
from .db import Database
from .export_xlsx import italian_date

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F6F3F")
BAND_FILL = PatternFill("solid", fgColor="EEF5F0")  # every other participant, for readability
MISSING_FILL = PatternFill("solid", fgColor="FFE699")  # bike still waiting for its number
NUMBER_FONT = Font(bold=True, size=14)
THIN = Side(style="thin", color="999999")
THICK = Side(style="medium", color="000000")
DATE_FORMAT = "DD/MM/YYYY"
TOP_ROWS = 3  # title, subtitle, blank; the header comes next

# (header, width, belongs to the participant -> merged over their bikes)
COLUMNS = [
    ("COGNOME", 16, True), ("NOME", 14, True), ("DATA DI NASCITA", 13, True), ("INDIRIZZO", 24, True),
    ("COMUNE", 15, True), ("PROV.", 7, True), ("TELEFONO", 14, True), ("EMAIL", 24, True),
    ("ASSICURATO", 13, True), ("N° POLIZZA", 14, True), ("PAGATO", 12, True),
    ("N°", 7, False), ("MARCA", 14, False), ("MODELLO", 22, False), ("CC", 7, False),
    ("ANNO", 7, False), ("CLASSE", 15, False), ("BATTERIA", 13, False), ("RITIRATA", 11, False),
    ("NOTE", 28, True),
]
LOOKUP_COLUMNS = [("N°", 7), ("MARCA", 14), ("MODELLO", 22), ("CC", 7), ("ANNO", 7), ("CLASSE", 15),
                  ("COGNOME", 16), ("NOME", 14)]


def build_export(db: Database, event_day: date | None = None) -> Workbook:
    riders = db.riders()  # already sorted by surname, name
    bikes_of = {r.id: sorted(db.bikes(r.id), key=lambda b: (b.race_number is None, b.race_number or 0))
                for r in riders}
    turn_of = {}
    for class_name, turns in db.grids().items():
        for t, rows in enumerate(turns, start=1):
            for row in rows:
                turn_of[row.bike_id] = f"{t}° turno" if len(turns) > 1 else "turno unico"

    wb = Workbook()
    ws = wb.active
    ws.title = "Iscritti"
    title = "MotoGP de Bosch – Elenco iscritti" + (f" – {italian_date(event_day)}" if event_day else "")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16)
    n_bikes = sum(len(b) for b in bikes_of.values())
    ws["A2"] = (f"{len(riders)} iscritti, {n_bikes} moto. Stampato il {date.today().strftime('%d/%m/%Y')}. "
                "Le caselle gialle sono moto ancora senza numero di gara.")
    ws["A2"].font = Font(italic=True, color="555555")
    header_row = TOP_ROWS + 1
    _header(ws, header_row, [(name, width) for name, width, _ in COLUMNS])

    row = header_row + 1
    for i, r in enumerate(riders):
        bikes = bikes_of[r.id] or [None]
        first, last = row, row + len(bikes) - 1
        rider_values = {
            "COGNOME": r.last_name, "NOME": r.first_name,
            "DATA DI NASCITA": date.fromisoformat(r.birth_date) if r.birth_date else None,
            "INDIRIZZO": r.address, "COMUNE": r.town, "PROV.": r.province, "TELEFONO": r.phone,
            "EMAIL": r.email, "ASSICURATO": "Sì" if r.insured else "No", "N° POLIZZA": r.insurance_id,
            "PAGATO": None,  # left empty on purpose: ticked by hand
            # NOTE is merged over all the participant's rows: collect every bike note here.
            "NOTE": "; ".join([*([r.notes] if r.notes else []), *(
                f"n° {b.race_number or '?'}: {b.notes}" for b in bikes_of[r.id] if b.notes)]),
        }
        for offset, b in enumerate(bikes):
            values = dict(rider_values) if offset == 0 else {}
            if b is not None:
                class_name = db.bike_class(b)
                values.update({
                    "N°": b.race_number, "MARCA": b.brand, "MODELLO": b.model, "CC": b.cc, "ANNO": b.year,
                    "CLASSE": CLASS_TITLES.get(class_name, class_name) if class_name else "",
                    "BATTERIA": turn_of.get(b.id, ""), "RITIRATA": "Sì" if b.withdrawn else "",
                })
            for c, (name, _, _) in enumerate(COLUMNS, start=1):
                cell = ws.cell(row + offset, c, values.get(name))
                cell.alignment = Alignment(vertical="center", wrap_text=True,
                                           horizontal="center" if name in ("N°", "CC", "ANNO", "PROV.", "ASSICURATO", "PAGATO", "RITIRATA") else None)
                if name == "DATA DI NASCITA":
                    cell.number_format = DATE_FORMAT
                if name == "N°":
                    cell.font = NUMBER_FONT
                    if b is not None and b.race_number is None:
                        cell.fill = MISSING_FILL
                elif i % 2:
                    cell.fill = BAND_FILL
                cell.border = Border(
                    left=THICK if name in ("PAGATO", "N°") else THIN,
                    right=THICK if name in ("PAGATO", "NOTE") or c == len(COLUMNS) else THIN,
                    top=THICK if offset == 0 else THIN,
                    bottom=THICK if row + offset == last else THIN,
                )
            ws.row_dimensions[row + offset].height = 22
        if last > first:
            for c, (name, _, merged) in enumerate(COLUMNS, start=1):
                if merged:
                    ws.merge_cells(start_row=first, start_column=c, end_row=last, end_column=c)
        row = last + 1

    ws.freeze_panes = ws.cell(header_row + 1, 3)
    _print_setup(ws, header_row)

    lookup = wb.create_sheet("Moto per numero")
    lookup["A1"] = "Moto in ordine di numero di gara"
    lookup["A1"].font = Font(bold=True, size=16)
    _header(lookup, TOP_ROWS + 1, LOOKUP_COLUMNS)
    names = {r.id: r for r in riders}
    all_bikes = sorted((b for bs in bikes_of.values() for b in bs),
                       key=lambda b: (b.race_number is None, b.race_number or 0))
    for b in all_bikes:
        class_name = db.bike_class(b)
        lookup.append([b.race_number, b.brand, b.model, b.cc, b.year,
                       CLASS_TITLES.get(class_name, class_name) if class_name else "",
                       names[b.rider_id].last_name, names[b.rider_id].first_name])
        number = lookup.cell(lookup.max_row, 1)
        number.font, number.alignment = NUMBER_FONT, Alignment(horizontal="center")
        if b.race_number is None:
            number.fill = MISSING_FILL
        for cell in lookup[lookup.max_row]:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    lookup.freeze_panes = lookup.cell(TOP_ROWS + 2, 1)
    _print_setup(lookup, TOP_ROWS + 1)
    return wb


def _header(ws, row: int, columns) -> None:
    for c, (name, width) in enumerate(columns, start=1):
        cell = ws.cell(row, c, name)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[row].height = 32


def _print_setup(ws, header_row: int) -> None:
    ws.print_title_rows = f"{header_row}:{header_row}"  # header repeated on every printed page
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.oddFooter.center.text = "Pagina &P di &N"
